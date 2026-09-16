# -*- coding: utf-8 -*-
"""
NeaT 공공급식 전자조달시스템(https://ns.eat.co.kr) 입찰공고 자동 수집 프로그램.

사용자가 수요기관명과 연도를 입력하면, 그 해 1월~12월 입찰공고를 한 달씩 조회해서
각 공고의 "기초가격"과 "낙찰예정가"를 가져와 엑셀 한 개로 합쳐준다.

[왜 한 달씩 조회하나]
입찰공고 화면의 "입찰기간" 조건은 1년을 통째로 넣어도 조회가 되지 않는다.
그래서 조회는 무조건 한 달 단위로 끊어서 열두 번 반복한다.

[대상월 vs 조회월]
9월분 급식 공고는 8월에 올라온다. 즉 "2025년 9월분" 데이터를 받으려면
입찰기간을 2025-08-01 ~ 2025-08-31 로 조회해야 한다.
이 프로그램에서 "대상월"은 급식월(9월), "조회월"은 그 전달(8월)을 뜻한다.

[넥사크로(nexacro) 사이트를 다루는 방법]
이 사이트는 화면 전체가 nexacro17로 그려져서 일반적인 <input>, <button> 태그가 없다.
DOM에는 "mainframe.VFS_MAIN...form.edt_instNm" 처럼 컴포넌트 경로가 그대로 id로 붙은
div만 있고, 실제 값은 전부 nexacro Dataset이 들고 있다.
그래서 화면 조작은 두 가지를 섞어서 한다.
 1) 사람이 눌러야만 화면이 열리는 것(메뉴, 공고번호 링크)은 div id로 click
 2) 조회조건 입력/결과 읽기는 nexacro JS API로 Dataset을 직접 읽고 쓴다
    (입력창에 한 글자씩 타이핑하는 것보다 훨씬 빠르고 안정적이다)
"""

import os
import re
import sys
import subprocess
from datetime import datetime

import openpyxl
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter

# 크로미움을 exe 안에 넣지 않으므로, 실행하는 PC의 기본 위치에서 찾도록 맞춰둔다.
# LOCALAPPDATA 는 윈도우에만 있는 값이라, 리눅스 서버에 올릴 때를 대비해 있을 때만 지정한다.
# (없으면 playwright 가 알아서 기본 위치를 쓴다. 도커처럼 밖에서 미리 지정해줬으면 그걸 존중한다.)
_LOCAL_APPDATA = os.environ.get("LOCALAPPDATA")
if _LOCAL_APPDATA and not os.environ.get("PLAYWRIGHT_BROWSERS_PATH"):
    os.environ["PLAYWRIGHT_BROWSERS_PATH"] = os.path.join(_LOCAL_APPDATA, "ms-playwright")

from playwright._impl._driver import compute_driver_executable, get_driver_env
from playwright.sync_api import sync_playwright


# ---------------------------------------------------------------------------
# 설정값
# ---------------------------------------------------------------------------

URL = "https://ns.eat.co.kr/NeaT/eats/index.html"

# 공고가 여러 건일 때 골라낼 품목 키워드. 공고건명에 "N월"과 이 키워드가 같이 들어간 건을 고른다.
# 수요기관마다 품목 표기가 달라서(공산품 / 가공식품류 ...) 실행할 때 바꿀 수 있게 해두었다.
DEFAULT_ITEM_KEYWORD = "공산품"

# nexacro 컴포넌트 경로. 화면 구조가 바뀌면 여기만 고치면 된다.
LEFT_FORM = "mainframe.VFS_MAIN.HFS_MAIN.CF_LEFT.form"
MDI_FORM = "mainframe.VFS_MAIN.HFS_MAIN.VFS_WORK.CF_MDI.form.div_mdi.form"
NOTICE_CLOSE_BTN = "mainframe.VFS_MAIN.CF_TOP.cmmNotice_PopList.form.div_popupTitle.form.btn_popClose"

# 메뉴번호(사이트 내부 메뉴 ID). 입찰공고=8060100, 입찰공고상세=8061000
MENU_BID_LIST = "8060100"
MENU_BID_DETAIL = "8061000"

WORK_FORM_FMT = "mainframe.VFS_MAIN.HFS_MAIN.VFS_WORK.FS_WORK.win{menu}.form.div_work.form"

# 넥사크로 앱이 완전히 뜰 때까지 꽤 오래 걸린다(스크립트 파일이 수백 개다).
APP_LOAD_TIMEOUT = 180_000
TRANSACTION_TIMEOUT = 60_000


class CollectError(Exception):
    """사용자에게 그대로 보여줄 수 있는 수준의 오류."""


def ensure_chromium_installed():
    # 크로미움이 아직 없으면 이 시점에 자동으로 받는다(인터넷 필요).
    # 이미 있으면 playwright가 알아서 바로 넘어간다.
    node_executable, cli_js = compute_driver_executable()
    subprocess.run(
        [str(node_executable), str(cli_js), "install", "chromium"],
        env=get_driver_env(),
        check=True,
    )


# ---------------------------------------------------------------------------
# nexacro 공용 헬퍼
# ---------------------------------------------------------------------------

def nx_eval(page, script, args=None, menu=MENU_BID_LIST):
    """업무화면(win{menu}) form을 f로 넘겨서 JS를 실행한다.

    nexacro는 모든 데이터를 Dataset에 들고 있기 때문에,
    f.ds_list.getColumn(...) 처럼 직접 읽는 게 화면 글자를 긁는 것보다 정확하다.
    """
    return page.evaluate(
        """([script, args, menu]) => {
            const app = nexacro.getApplication();
            const win = app.mainframe.VFS_MAIN.HFS_MAIN.VFS_WORK.FS_WORK["win" + menu];
            if (!win) throw new Error("화면(win" + menu + ")이 열려있지 않습니다.");
            const f = win.form.div_work.form;
            return (new Function("f", "args", "app", script))(f, args, app);
        }""",
        [script, args, menu],
    )


def nx_dataset_rows(page, ds_name, columns, menu=MENU_BID_LIST):
    """Dataset을 파이썬 dict 리스트로 읽어온다.

    Dataset 값에는 날짜형이 섞여 있어서 그대로 넘기면 JSON 직렬화가 깨진다.
    그래서 JS 쪽에서 전부 문자열로 바꿔서 넘긴다.
    """
    return nx_eval(
        page,
        """
        const [dsName, cols] = args;
        const ds = f[dsName];
        if (!ds) return [];
        const out = [];
        for (let r = 0; r < ds.getRowCount(); r++) {
            const row = {};
            for (const c of cols) {
                const v = ds.getColumn(r, c);
                row[c] = (v === null || v === undefined) ? "" : String(v);
            }
            out.push(row);
        }
        return out;
        """,
        [ds_name, columns],
        menu,
    )


def click_component(page, component_id, timeout=20_000):
    """div id가 nexacro 컴포넌트 경로인 요소를 클릭한다."""
    page.locator('div[id="%s"]' % component_id).first.click(timeout=timeout)


# ---------------------------------------------------------------------------
# 1) 사이트 접속 ~ 입찰공고 화면 열기
# ---------------------------------------------------------------------------

def open_site(playwright, headless=True):
    # 리눅스 컨테이너(도커/Render)에서는 이 두 옵션이 없으면 크로미움이 바로 죽는다.
    #  --no-sandbox          : 컨테이너는 보통 root로 도는데, 그러면 샌드박스가 기동을 거부한다
    #  --disable-dev-shm-usage : 컨테이너의 /dev/shm(기본 64MB)이 작아서 탭이 뻗는다
    launch_args = [] if os.name == "nt" else ["--no-sandbox", "--disable-dev-shm-usage"]

    browser = playwright.chromium.launch(headless=headless, args=launch_args)
    page = browser.new_page(viewport={"width": 1680, "height": 980})

    # networkidle로 기다리면 넥사크로가 계속 통신을 해서 타임아웃이 난다.
    # 문서만 먼저 받고, 앱 객체가 만들어질 때까지 따로 기다린다.
    page.goto(URL, wait_until="domcontentloaded", timeout=APP_LOAD_TIMEOUT)
    page.wait_for_function(
        """() => {
            try {
                return !!(nexacro.getApplication && nexacro.getApplication())
                    && document.querySelectorAll('div[id^="mainframe"]').length > 5;
            } catch (e) { return false; }
        }""",
        timeout=APP_LOAD_TIMEOUT,
    )
    page.wait_for_timeout(3000)

    close_notice_popup(page)
    return browser, page


def close_notice_popup(page):
    """첫 화면에 뜨는 "팝업공지사항"을 닫는다.

    이 팝업은 모달이라서 닫지 않으면 뒤에 있는 햄버거 버튼 클릭이 막힌다.
    공지가 없는 날에는 팝업 자체가 없으므로 없으면 그냥 넘어간다.
    """
    popup = page.locator('div[id="%s"]' % NOTICE_CLOSE_BTN)
    if popup.count() == 0:
        return
    try:
        popup.first.click(timeout=5000)
        page.wait_for_timeout(1500)
    except Exception:
        pass


def _visible_menu_rows(page):
    """왼쪽 메뉴(트리 그리드)에서 지금 화면에 보이는 항목들을 읽는다.

    메뉴 그리드가 div_leftMenu / pdiv_leftMenu 두 벌로 존재하고
    둘 다 DOM에는 남아있기 때문에, 반드시 "보이는" 쪽만 골라야 클릭이 된다.
    """
    return page.evaluate(
        """() => Array.from(
            document.querySelectorAll('div[id*="grd_leftMenu.body"][id$="treeitemtext:text"]'))
            .filter(d => d.offsetWidth > 0 && d.offsetHeight > 0)
            .map(d => [d.id, (d.innerText || "").trim()])"""
    )


def _click_menu(page, menu_name):
    for component_id, text in _visible_menu_rows(page):
        if text == menu_name:
            click_component(page, component_id)
            return True
    return False


def open_bid_list_screen(page):
    """햄버거 버튼 -> 입찰정보 -> 입찰공고 순서로 눌러 업무화면을 연다."""
    click_component(page, "%s.btn_8000000:icontext" % LEFT_FORM)
    page.wait_for_timeout(2000)

    if not _click_menu(page, "입찰정보"):
        raise CollectError("왼쪽 메뉴에서 '입찰정보'를 찾지 못했습니다.")
    page.wait_for_timeout(2000)

    if not _click_menu(page, "입찰공고"):
        raise CollectError("왼쪽 메뉴에서 '입찰공고'를 찾지 못했습니다.")

    # 조회조건의 수요기관명 입력창이 그려지면 화면이 다 뜬 것으로 본다.
    page.wait_for_function(
        """() => !!document.getElementById(
            "%s.div_search.form.edt_instNm")""" % WORK_FORM_FMT.format(menu=MENU_BID_LIST),
        timeout=APP_LOAD_TIMEOUT,
    )
    page.wait_for_timeout(2000)


# ---------------------------------------------------------------------------
# 2) 조회
# ---------------------------------------------------------------------------

# 조회 결과(ds_list)에서 실제로 쓰는 컬럼만 추린다.
LIST_COLUMNS = [
    "ETN_BID_ID",    # 내부 입찰 ID
    "ETN_BID_NO",    # 공고번호
    "ETN_BID_STT_NM",  # 진행상태(진행중/낙찰/유찰 등)
    "PURR_NM",       # 수요기관명
    "BID_NM",        # 공고건명
    "PBANC_YMD",     # 공고개시일
    "STRPRCE",       # 기초가격(목록에도 있지만 확인용)
    "PLNPRCE",       # 낙찰예정가(진행중이면 0)
]


def search_month(page, inst_nm, bgn_ymd, end_ymd):
    """수요기관명 + 입찰기간으로 조회하고 결과 목록을 돌려준다.

    조회조건 컴포넌트들은 전부 ds_searchParam에 바인딩되어 있어서
    Dataset만 채우고 fn_search()를 부르면 화면에 타이핑한 것과 똑같이 동작한다.

    조회는 비동기(서버 트랜잭션)라서 끝나는 시점을 알아야 하는데,
    nexacro는 응답이 오면 화면의 fn_callback을 부른다.
    그래서 fn_callback을 한 번만 감싸두고(__auto_done 플래그) 그게 켜질 때까지 기다린다.
    """
    nx_eval(
        page,
        """
        const [instNm, bgn, end] = args;

        // 조회 완료를 감지하기 위해 fn_callback을 감싼다(중복 래핑 방지).
        if (!f.__auto_wrapped) {
            const orig = f.fn_callback;
            f.fn_callback = function (svcId, errCd, errMsg) {
                f.__auto_done = true;
                return orig.call(f, svcId, errCd, errMsg);
            };
            f.__auto_wrapped = true;
        }
        f.__auto_done = false;

        const ds = f.ds_searchParam;
        ds.setColumn(0, "P_INST_NM", instNm);
        ds.setColumn(0, "P_BID_BGNG_DT", bgn);
        ds.setColumn(0, "P_BID_END_DT", end);
        ds.setColumn(0, "P_BID_NM", "");
        ds.setColumn(0, "P_ELCTRN_BID_NO", "");

        f.ds_list.clearData();
        f.fn_search();
        return true;
        """,
        [inst_nm, bgn_ymd, end_ymd],
    )

    page.wait_for_function(
        """() => {
            const app = nexacro.getApplication();
            const f = app.mainframe.VFS_MAIN.HFS_MAIN.VFS_WORK.FS_WORK.win%s.form.div_work.form;
            return f.__auto_done === true;
        }""" % MENU_BID_LIST,
        timeout=TRANSACTION_TIMEOUT,
    )
    page.wait_for_timeout(1000)

    return nx_dataset_rows(page, "ds_list", LIST_COLUMNS)


def months_in_title(title):
    """공고건명에 적힌 급식월을 모두 뽑는다.

    "2026년 9월", "2025학년도 10월분" 처럼 한 달짜리가 대부분이지만,
    방학이 낀 달은 "2026년 7~8월" 처럼 두 달을 묶어서 한 번만 공고한다.
    그래서 범위 표기를 먼저 풀어준 뒤 단일 표기를 더한다.
    """
    months = set()

    for match in re.finditer(r"(\d{1,2})\s*[~∼\-–·,]\s*(\d{1,2})\s*월", title):
        first, second = int(match.group(1)), int(match.group(2))
        if 1 <= first <= 12 and 1 <= second <= 12:
            months.update(range(min(first, second), max(first, second) + 1))

    for match in re.finditer(r"(?<!\d)(\d{1,2})\s*월", title):
        value = int(match.group(1))
        if 1 <= value <= 12:
            months.add(value)

    return months


def pick_target_row(rows, target_month, item_keyword=DEFAULT_ITEM_KEYWORD, choose_func=None):
    """한 달 조회 결과에서 가져올 공고 한 건을 고른다.

    기본적으로 결과는 한 건이지만, 여러 건이면 공고건명에
    대상월("9월")과 품목 키워드("공산품")가 동시에 들어있는 건을 고른다.
    그래도 여러 건이 남거나 키워드가 하나도 안 맞으면
    choose_func으로 사용자에게 직접 고르게 한다.
    """
    if not rows:
        return None, "조회결과 없음"

    if len(rows) == 1:
        row = rows[0]
        months = months_in_title(row["BID_NM"])
        if months and target_month not in months:
            return row, "조회결과가 1건이지만 공고건명의 급식월(%s)이 대상월과 다름" % \
                ", ".join("%d월" % m for m in sorted(months))
        return row, ""

    month_rows = [r for r in rows if target_month in months_in_title(r["BID_NM"])]
    if not month_rows:
        return None, "공고건명에 '%d월'이 들어간 건이 없음(총 %d건)" % (target_month, len(rows))

    matched = [r for r in month_rows if item_keyword in r["BID_NM"]] if item_keyword else month_rows

    if len(matched) == 1:
        return matched[0], ""

    if not matched:
        # 수요기관마다 품목 이름이 달라서(공산품 / 가공식품류 ...) 키워드가 안 맞을 수 있다.
        # 이때는 그냥 넘기지 말고 그 달의 공고를 전부 보여주고 고르게 한다.
        note = "'%s' 포함 공고가 없어 직접 선택" % item_keyword
        candidates = month_rows
    else:
        note = ""
        candidates = matched

    if choose_func is None:
        return candidates[0], (note + " / " if note else "") + \
            "조건에 맞는 건이 %d건이라 첫 번째 건을 사용" % len(candidates)

    chosen = choose_func(candidates)
    if chosen is None:
        return None, (note + " / " if note else "") + "사용자가 선택하지 않음"
    return chosen, note


# ---------------------------------------------------------------------------
# 3) 입찰공고상세에서 기초가격 / 낙찰예정가 읽기
# ---------------------------------------------------------------------------

def read_detail(page, bid_no):
    """목록에서 공고번호를 클릭해 상세화면을 열고, 금액을 읽은 뒤 다시 닫는다.

    상세화면은 새 MDI 탭(win8061000)으로 열린다.
    금액은 화면 글자가 아니라 ds_info에서 읽는다.
      - BGNG_PRC          : 기초가격
      - ELCTRN_BID_PLNPRC : 낙찰예정가격
      - PLNPRCE_SUCBD_STD : 낙찰하한율(%) ("예정가격의 [88]%이상 ... 최저가 낙찰")
    """
    cell_id = _find_bid_no_cell(page, bid_no)
    if cell_id is None:
        raise CollectError("목록에서 공고번호 %s 를 찾지 못했습니다." % bid_no)

    click_component(page, cell_id)

    # 상세화면 form이 만들어지고 ds_info에 데이터가 채워질 때까지 기다린다.
    page.wait_for_function(
        """() => {
            try {
                const app = nexacro.getApplication();
                const win = app.mainframe.VFS_MAIN.HFS_MAIN.VFS_WORK.FS_WORK.win%s;
                if (!win || !win.form || !win.form.div_work) return false;
                const f = win.form.div_work.form;
                return !!f.ds_info && f.ds_info.getRowCount() > 0;
            } catch (e) { return false; }
        }""" % MENU_BID_DETAIL,
        timeout=TRANSACTION_TIMEOUT,
    )
    page.wait_for_timeout(1500)

    info = nx_eval(
        page,
        """
        const ds = f.ds_info;
        const get = (c) => {
            const v = ds.getColumn(0, c);
            return (v === null || v === undefined) ? "" : String(v);
        };
        return {
            bid_no: get("ELCTRN_BID_NO"),
            bid_nm: get("BID_NM"),
            base_price: get("BGNG_PRC"),
            plan_price: get("ELCTRN_BID_PLNPRC"),
            lower_rate: get("PLNPRCE_SUCBD_STD"),
            status_nm: get("ETN_BID_STT_NM"),
        };
        """,
        None,
        MENU_BID_DETAIL,
    )

    close_detail_screen(page)
    return info


def _find_bid_no_cell(page, bid_no):
    """그리드에서 공고번호 글자가 일치하는 셀의 div id를 찾는다.

    행 번호(gridrow_N)로 찾으면 정렬/페이징에 따라 어긋날 수 있어서,
    화면에 찍힌 공고번호 글자로 직접 찾는다.
    """
    cells = page.evaluate(
        """(gridId) => Array.from(document.querySelectorAll('div[id^="' + gridId + '.body.gridrow_"]'))
            .filter(d => d.id.endsWith(":text") && d.offsetWidth > 0)
            .map(d => [d.id, (d.innerText || "").trim()])""",
        "%s.grd_list" % WORK_FORM_FMT.format(menu=MENU_BID_LIST),
    )
    for component_id, text in cells:
        if text == bid_no:
            return component_id
    return None


def close_detail_screen(page):
    """상세 탭을 닫아 목록 화면으로 돌아온다.

    탭을 계속 열어두면 다음 달 조회 때 상세 탭이 활성 상태로 남아
    목록 그리드 클릭이 안 된다.
    """
    close_btn = "%s.btn_mdiClose%s" % (MDI_FORM, "win" + MENU_BID_DETAIL)
    locator = page.locator('div[id="%s"]' % close_btn)
    if locator.count():
        locator.first.click(timeout=15_000)

    page.wait_for_function(
        """() => {
            try {
                const app = nexacro.getApplication();
                return !app.mainframe.VFS_MAIN.HFS_MAIN.VFS_WORK.FS_WORK.win%s;
            } catch (e) { return false; }
        }""" % MENU_BID_DETAIL,
        timeout=30_000,
    )
    page.wait_for_timeout(1000)


# ---------------------------------------------------------------------------
# 4) 기간 계산
# ---------------------------------------------------------------------------

def add_month(year, month, diff):
    idx = (year * 12 + (month - 1)) + diff
    return idx // 12, idx % 12 + 1


def last_day(year, month):
    next_year, next_month = add_month(year, month, 1)
    return (datetime(next_year, next_month, 1) - datetime(year, month, 1)).days


def make_period(year, month):
    """대상월(급식월) 하나에 대한 조회기간을 만든다.

    대상월 2025-09 -> 조회기간 2025-08-01 ~ 2025-08-31 (공고는 한 달 전에 올라오므로)
    """
    s_year, s_month = add_month(year, month, -1)
    return {
        "target_year": year,
        "target_month": month,
        "label": "%04d-%02d" % (year, month),
        "bgn": "%04d%02d01" % (s_year, s_month),
        "end": "%04d%02d%02d" % (s_year, s_month, last_day(s_year, s_month)),
    }


def build_periods_for_year(year, today=None):
    """그 해 1월~12월의 대상월 목록을 만든다.

    아직 오지 않은 달은 공고 자체가 없으므로 빼고,
    이번 달까지만 조회한다(조회 한 번에 5초씩 걸려서 헛걸음을 줄인다).
    """
    today = today or datetime.today()

    periods = []
    for month in range(1, 13):
        if (year, month) > (today.year, today.month):
            break
        periods.append(make_period(year, month))
    return periods


# ---------------------------------------------------------------------------
# 5) 수집 본체
# ---------------------------------------------------------------------------

def to_int(text):
    if text is None:
        return None
    digits = re.sub(r"[^\d]", "", str(text))
    return int(digits) if digits else None


def find_in_previous_rows(previous_rows, target_month, item_keyword=DEFAULT_ITEM_KEYWORD):
    """직전 달 조회결과에서 대상월을 함께 포함하는 공고를 찾는다.

    방학이 낀 달은 "2026년 7~8월"처럼 두 달을 한 번에 공고해버려서,
    8월분으로 조회해야 할 7월에는 아무 공고도 올라오지 않는다.
    이럴 때 직전 달(6월) 조회결과에 이미 받아둔 통합공고를 그대로 쓴다.
    """
    for row in previous_rows:
        if target_month in months_in_title(row["BID_NM"]) and item_keyword in row["BID_NM"]:
            return row
    return None


def collect(page, inst_nm, periods, log=print, choose_func=None,
            item_keyword=DEFAULT_ITEM_KEYWORD, stop_check=None, on_record=None):
    """월별로 조회 -> 공고 선택 -> 상세 금액 읽기를 반복한다.

    stop_check : 매달 시작 전에 불러서 True면 중단한다(UI의 "중지" 버튼용).
    on_record  : 한 달이 끝날 때마다 그 결과를 넘겨준다(UI 진행률/목록 갱신용).
    """
    results = []
    previous_rows = []
    previous_range = None

    for period in periods:
        if stop_check is not None and stop_check():
            log("중지 요청으로 수집을 멈췄습니다.")
            break

        log("[%s] 조회 중... (입찰기간 %s ~ %s)" % (period["label"], period["bgn"], period["end"]))

        record = {
            "label": period["label"],
            "target_year": period["target_year"],
            "target_month": period["target_month"],
            "bgn": period["bgn"],
            "end": period["end"],
            "bid_no": "",
            "bid_nm": "",
            "pbanc_ymd": "",
            "status": "",
            "base_price": None,
            "plan_price": None,
            "lower_rate": None,
            "note": "",
        }

        try:
            rows = search_month(page, inst_nm, period["bgn"], period["end"])
        except Exception as error:
            record["note"] = "조회 실패: %s" % error
            log("  -> %s" % record["note"])
            _add_record(results, record, on_record)
            previous_rows, previous_range = [], None
            continue

        row, note = pick_target_row(rows, period["target_month"], item_keyword, choose_func)
        detail_range = (period["bgn"], period["end"])

        if row is None:
            # 두 달을 묶은 통합공고(예: "7~8월")라서 이번 달에는 공고가 없는 경우
            fallback = find_in_previous_rows(previous_rows, period["target_month"], item_keyword)
            if fallback is not None:
                row = fallback
                note = "직전 달에 올라온 통합공고 사용"
                # 상세화면은 목록 그리드의 공고번호를 눌러서 여는 구조라,
                # 그 공고가 화면에 떠 있도록 직전 달 조건으로 다시 조회한다.
                detail_range = previous_range
                # 비고와 조회기간이 서로 안 맞으면 나중에 헷갈리므로 실제 찾은 기간으로 남긴다.
                record["bgn"], record["end"] = previous_range

        previous_rows, previous_range = rows, (period["bgn"], period["end"])

        if row is None:
            record["note"] = note
            log("  -> %s" % note)
            _add_record(results, record, on_record)
            continue

        record["bid_no"] = row["ETN_BID_NO"]
        record["bid_nm"] = row["BID_NM"]
        record["pbanc_ymd"] = row["PBANC_YMD"]
        record["status"] = row["ETN_BID_STT_NM"]
        record["note"] = note

        try:
            if detail_range != (period["bgn"], period["end"]):
                search_month(page, inst_nm, detail_range[0], detail_range[1])
            detail = read_detail(page, row["ETN_BID_NO"])
        except Exception as error:
            record["note"] = (record["note"] + " / " if record["note"] else "") + \
                "상세조회 실패: %s" % error
            log("  -> %s" % record["note"])
            _add_record(results, record, on_record)
            continue

        record["base_price"] = to_int(detail["base_price"])
        record["plan_price"] = to_int(detail["plan_price"])
        record["lower_rate"] = to_int(detail["lower_rate"])

        # 아직 개찰 전(진행중)이면 낙찰예정가가 0으로 내려온다.
        if not record["plan_price"]:
            record["plan_price"] = None
            record["note"] = (record["note"] + " / " if record["note"] else "") + \
                "낙찰예정가 미공개(진행상태: %s)" % record["status"]

        log("  -> %s / 기초가격 %s / 낙찰예정가 %s" % (
            record["bid_no"],
            format(record["base_price"], ",") if record["base_price"] else "-",
            format(record["plan_price"], ",") if record["plan_price"] else "-",
        ))
        _add_record(results, record, on_record)

    return results


def _add_record(results, record, on_record):
    results.append(record)
    if on_record is not None:
        on_record(record)


# ---------------------------------------------------------------------------
# 6) 엑셀 저장
# ---------------------------------------------------------------------------

# 투찰.xlsx 가 쓰고 있는 서식을 그대로 가져왔다(회계 표시형식 / 투찰률 소수 5자리).
ACC_INT = r'_-* #,##0_-;\-* #,##0_-;_-* "-"_-;_-@_-'
ACC_DEC2 = r'_-* #,##0.00_-;\-* #,##0.00_-;_-* "-"??_-;_-@_-'
ACC_DEC5 = r'_-* #,##0.00000_-;\-* #,##0.00000_-;_-* "-"?????_-;_-@_-'
RATE_FMT = "0.00000_ "
COUNT_FMT = "#,##0_ "

BASE_FONT = Font(name="맑은 고딕", size=11)
CENTER = Alignment(horizontal="center", vertical="center")

# 투찰.xlsx 가 "플러스/마이너스" 칸에 걸어둔 조건부서식 색상(엑셀 기본 좋음/나쁨 색).
PLUS_STYLE = DifferentialStyle(font=Font(color="FF006100"), fill=PatternFill(bgColor="FFC6EFCE"))
MINUS_STYLE = DifferentialStyle(font=Font(color="FF9C0006"), fill=PatternFill(bgColor="FFFFC7CE"))

# 낙찰하한율을 못 읽었을 때 쓸 기본값(투찰.xlsx 에서 가장 많이 쓰는 값)
DEFAULT_LOWER_RATE = 88


def lower_rate_of(record):
    return record["lower_rate"] or DEFAULT_LOWER_RATE


def bid_rate(record):
    """투찰률 = 예정가 / 기초가 - (1 - 낙찰하한율).

    투찰.xlsx 의 투찰률 수식(=D3/C3-0.12)과 같은 계산이다.
    빼는 값 0.12 / 0.1 은 낙찰하한율 88% / 90% 에서 나온 것이라,
    공고 상세에서 읽어온 하한율로 자동으로 맞춰준다.
    """
    if not record["base_price"] or not record["plan_price"]:
        return None
    return record["plan_price"] / record["base_price"] - (1 - lower_rate_of(record) / 100.0)


def save_excel(inst_nm, records, out_dir):
    workbook = openpyxl.Workbook()

    _write_format_sheet(workbook.active, inst_nm, records)
    _write_detail_sheet(workbook.create_sheet("상세"), inst_nm, records)

    first, last = records[0]["label"], records[-1]["label"]
    base_name = "%s_입찰데이터_%s~%s" % (inst_nm, first, last)

    # 같은 이름의 파일을 엑셀로 열어둔 채 다시 돌리면 덮어쓸 수가 없다(PermissionError).
    # 애써 모은 데이터를 날리지 않도록, 뒤에 번호를 붙여서라도 저장한다.
    for suffix in [""] + ["(%d)" % n for n in range(2, 21)]:
        file_path = os.path.join(out_dir, base_name + suffix + ".xlsx")
        try:
            workbook.save(file_path)
            return file_path
        except PermissionError:
            continue

    raise CollectError("엑셀 파일을 저장하지 못했습니다. 열려있는 엑셀 파일을 닫고 다시 실행해주세요.")


def _sheet_title(inst_nm):
    """시트명은 수요기관명으로 한다(엑셀에서 못 쓰는 문자만 빼고).

    한 시트 안에 연도별 블록이 쌓이는 구조라, 시트명에는 연도를 넣지 않는다.
    """
    title = inst_nm
    for char in "[]:*?/\\":
        title = title.replace(char, "")
    return title[:31] or "수집결과"


def _year_label(records):
    """1행 A칸에 들어갈 연도 표기. 두 해에 걸치면 "25~26년" 으로 쓴다."""
    years = sorted({record["target_year"] for record in records})
    if len(years) == 1:
        return "%02d년" % (years[0] % 100)
    return "%02d~%02d년" % (years[0] % 100, years[-1] % 100)


def _write_format_sheet(sheet, inst_nm, records):
    """투찰.xlsx 와 똑같은 양식으로 쓴다.

      1행 : [26년][   ] [1월(3칸 병합)] [2월] ... [12월]
      2행 : [구분 ][평균] [기초가][예정가][투찰률] ...
      3행 : [기관명][평균수식] [기초가][예정가][=예정가/기초가-0.12] ...
      4행 : [ ↑병합][건수수식] [=기초가*투찰률][=투찰금액/예정가][=IF(투찰률<0.88,"마이너스","플러스")]

    월 칸은 데이터가 있든 없든 1월~12월 12칸을 항상 만든다.
    그래야 어느 기관을 뽑아도 "3월은 항상 H열" 처럼 열 위치가 고정된다.

    최근 12개월을 모으면 연도가 두 해에 걸치므로(예: 25년 10월 ~ 26년 9월),
    투찰.xlsx 가 "능실초(25년)/능실초(26년)" 처럼 나눠 쓰는 방식 그대로
    연도마다 2행짜리 블록을 따로 만든다.

    계산은 전부 파이썬이 아니라 엑셀 수식으로 넣는다.
    그래야 담당자가 기초가·예정가를 손으로 고치면 투찰률이 같이 따라 움직인다.
    """
    sheet.title = _sheet_title(inst_nm)

    last_col = 2 + 12 * 3
    last_letter = get_column_letter(last_col)

    # 1~2행 머리글
    sheet.cell(row=1, column=1, value=_year_label(records))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    sheet.cell(row=2, column=1, value="구분")
    sheet.cell(row=2, column=2, value="평균")

    for month in range(1, 13):
        col = 3 + (month - 1) * 3
        sheet.cell(row=1, column=col, value="%d월" % month)
        sheet.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
        sheet.cell(row=2, column=col, value="기초가")
        sheet.cell(row=2, column=col + 1, value="예정가")
        sheet.cell(row=2, column=col + 2, value="투찰률")

    # 연도별로 묶어서 오래된 연도부터 2행씩 채운다.
    years = sorted({record["target_year"] for record in records})
    row = 3
    for year in years:
        by_month = {r["target_month"]: r for r in records if r["target_year"] == year}
        label = inst_nm if len(years) == 1 else "%s(%02d년)" % (inst_nm, year % 100)
        _write_year_block(sheet, row, label, by_month, last_letter)
        row += 2

    _apply_format_style(sheet, last_col, row - 1)


def _write_year_block(sheet, row, label, by_month, last_letter):
    """한 기관·한 연도를 2행으로 쓴다."""
    row2 = row + 1

    sheet.cell(row=row, column=1, value=label)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row2, end_column=1)

    average = sheet.cell(row=row, column=2,
                         value='=(SUMIF(C%d:%s%d,"<1"))/B%d' % (row, last_letter, row, row2))
    average.number_format = ACC_DEC5

    count = sheet.cell(row=row2, column=2,
                       value='=COUNTIF(C%d:%s%d,"마이너스")+COUNTIF(C%d:%s%d,"플러스")'
                             % (row, last_letter, row2, row, last_letter, row2))
    count.number_format = COUNT_FMT

    for month in range(1, 13):
        col = 3 + (month - 1) * 3
        base_letter = get_column_letter(col)
        plan_letter = get_column_letter(col + 1)
        rate_letter = get_column_letter(col + 2)

        record = by_month.get(month)

        # 낙찰하한율 88% -> 수식에서 0.12 를 빼고, 마이너스/플러스 기준도 0.88 이 된다.
        lower = lower_rate_of(record) / 100.0 if record else DEFAULT_LOWER_RATE / 100.0
        margin = "%g" % round(1 - lower, 4)
        limit = "%g" % round(lower, 4)

        base_cell = sheet.cell(row=row, column=col,
                               value=record["base_price"] if record else None)
        base_cell.number_format = ACC_INT
        plan_cell = sheet.cell(row=row, column=col + 1,
                               value=record["plan_price"] if record else None)
        plan_cell.number_format = ACC_INT

        # 데이터가 없는 달도 수식은 그대로 넣는다(투찰.xlsx 도 빈 달은 #DIV/0! 로 남아있다).
        rate_cell = sheet.cell(row=row, column=col + 2,
                               value="=%s%d/%s%d-%s" % (plan_letter, row, base_letter, row, margin))
        rate_cell.number_format = RATE_FMT

        amount_cell = sheet.cell(row=row2, column=col,
                                 value="=%s%d*%s%d" % (base_letter, row, rate_letter, row))
        amount_cell.number_format = ACC_DEC2

        sheet.cell(row=row2, column=col + 1,
                   value="=%s%d/%s%d" % (base_letter, row2, plan_letter, row))

        sheet.cell(row=row2, column=col + 2,
                   value='=IF(%s%d<%s,"마이너스","플러스")' % (rate_letter, row, limit))


def _apply_format_style(sheet, last_col, last_row):
    """글꼴·가운데정렬·열너비를 투찰.xlsx 와 맞춘다."""
    for row in range(1, last_row + 1):
        for col in range(1, last_col + 1):
            cell = sheet.cell(row=row, column=col)
            cell.font = BASE_FONT
            # 머리글(1~2행), 구분·평균(A·B열), 투찰률·판정(3칸 중 마지막)만 가운데 정렬.
            if row <= 2 or col <= 2 or (col >= 3 and (col - 3) % 3 == 2):
                cell.alignment = CENTER

    sheet.column_dimensions["A"].width = 13.0
    sheet.column_dimensions["B"].width = 10.25
    for col in range(3, last_col + 1):
        width = (15.38, 14.25, 9.62)[(col - 3) % 3]
        sheet.column_dimensions[get_column_letter(col)].width = width

    sheet.freeze_panes = "C3"

    _apply_plus_minus_color(sheet, "A1:%s%d" % (get_column_letter(last_col), last_row))


def _apply_plus_minus_color(sheet, cell_range):
    """"플러스"는 초록, "마이너스"는 빨강으로 칠한다.

    셀마다 직접 색을 넣지 않고 투찰.xlsx 와 똑같이 조건부서식으로 건다.
    그래야 담당자가 기초가·예정가를 고쳐서 판정이 뒤집히면 색도 같이 바뀐다.
    """
    first_cell = cell_range.split(":")[0]

    for text, style in (("플러스", PLUS_STYLE), ("마이너스", MINUS_STYLE)):
        rule = Rule(type="containsText", operator="containsText", text=text, dxf=style)
        rule.formula = ['NOT(ISERROR(SEARCH("%s",%s)))' % (text, first_cell)]
        sheet.conditional_formatting.add(cell_range, rule)


def _write_detail_sheet(sheet, inst_nm, records):
    """월별 원본 데이터(공고번호/공고건명/금액/비고)를 그대로 남긴다."""
    headers = [
        "대상월", "수요기관명", "공고번호", "공고건명", "공고개시일", "진행상태",
        "기초가격", "낙찰예정가", "낙찰하한율(%)", "투찰률", "조회기간", "비고",
    ]
    sheet.append(headers)

    for record in records:
        rate = bid_rate(record)
        sheet.append([
            record["label"],
            inst_nm,
            record["bid_no"],
            record["bid_nm"],
            _format_ymd(record["pbanc_ymd"]),
            record["status"],
            record["base_price"],
            record["plan_price"],
            record["lower_rate"],
            rate,
            "%s ~ %s" % (_format_ymd(record["bgn"]), _format_ymd(record["end"])),
            record["note"],
        ])

    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(name="맑은 고딕", size=11, bold=True)
        cell.alignment = CENTER

    for row in range(2, len(records) + 2):
        for col in range(1, len(headers) + 1):
            sheet.cell(row=row, column=col).font = BASE_FONT
        sheet.cell(row=row, column=7).number_format = ACC_INT
        sheet.cell(row=row, column=8).number_format = ACC_INT
        sheet.cell(row=row, column=10).number_format = RATE_FMT

    widths = [10, 18, 18, 55, 12, 10, 14, 14, 13, 10, 25, 40]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _format_ymd(value):
    text = str(value or "")
    if len(text) >= 8 and text[:8].isdigit():
        return "%s-%s-%s" % (text[:4], text[4:6], text[6:8])
    return text


# ---------------------------------------------------------------------------
# 7) 콘솔 입력
# ---------------------------------------------------------------------------

def input_inst_nm():
    while True:
        value = input("수요기관명을 정확하게 입력하세요 (예: 기산중학교): ").strip()
        if value:
            return value
        print("수요기관명을 입력해주세요.\n")


def input_year():
    """수집할 연도(한 해). 그냥 엔터면 올해."""
    today = datetime.today()

    while True:
        value = input("수집할 연도를 입력하세요 (예: 2026) [기본값 %d]: " % today.year).strip()
        if not value:
            return today.year

        if re.fullmatch(r"\d{4}", value):
            return int(value)

        print("연도를 네 자리 숫자로 입력해주세요.\n")


def input_item_keyword():
    """공고가 여러 건일 때 고를 품목 키워드. 그냥 엔터면 "공산품"."""
    value = input("품목 키워드를 입력하세요 [기본값 %s]: " % DEFAULT_ITEM_KEYWORD).strip()
    return value or DEFAULT_ITEM_KEYWORD


def choose_from_console(rows):
    """공고가 여러 건 남았을 때 콘솔에서 고르게 한다."""
    print("  조건에 맞는 공고가 %d건입니다. 번호를 선택하세요." % len(rows))
    for index, row in enumerate(rows, start=1):
        print("   %d. %s | %s" % (index, row["ETN_BID_NO"], row["BID_NM"]))

    choice = input("  번호를 입력하세요 (건너뛰려면 엔터): ").strip()
    if choice.isdigit() and 1 <= int(choice) <= len(rows):
        return rows[int(choice) - 1]
    return None


def main():
    print("=" * 70)
    print(" NeaT 공공급식 입찰공고 수집 프로그램")
    print("=" * 70)

    inst_nm = input_inst_nm()
    year = input_year()
    item_keyword = input_item_keyword()
    periods = build_periods_for_year(year)

    if not periods:
        print("\n조회할 달이 없습니다. 연도를 확인해주세요.")
        return

    print("\n수요기관명 : %s" % inst_nm)
    print("수집연도   : %d년" % year)
    print("대상기간   : %s ~ %s (%d개월)" % (periods[0]["label"], periods[-1]["label"], len(periods)))
    print("품목 키워드 : %s\n" % item_keyword)

    print("크로미움 확인 중...")
    ensure_chromium_installed()

    with sync_playwright() as playwright:
        browser, page = open_site(playwright, headless=True)
        try:
            print("입찰공고 화면 여는 중...")
            open_bid_list_screen(page)
            print("입찰공고 화면 이동 완료\n")

            records = collect(page, inst_nm, periods,
                              choose_func=choose_from_console, item_keyword=item_keyword)
        finally:
            browser.close()

    found = [r for r in records if r["base_price"]]
    if not found:
        print("\n수집된 데이터가 없습니다. 수요기관명이 정확한지 확인해주세요.")
        return

    out_dir = os.path.dirname(os.path.abspath(__file__))
    file_path = save_excel(inst_nm, records, out_dir)

    print("\n%d개월 중 %d개월 수집 완료" % (len(records), len(found)))
    print("엑셀 저장 : %s" % file_path)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n사용자가 중단했습니다.")
        sys.exit(1)
    except CollectError as error:
        print("\n[오류] %s" % error)
        sys.exit(1)
